"""Extract text from course materials (PDF + PPTX).

PPTX is preferred when available — it's structured XML, gives us clean
text per shape, and exposes speaker notes (where the real substance often
lives). PDF is layout-guessed via PyMuPDF.

CS slides frequently embed code as screenshots (syntax-highlighted
images) and explanatory diagrams as raster pictures. Pure text extraction
misses those entirely, so we OCR picture shapes (PPTX) and pages
containing images (PDF) via Tesseract. OCR'd content is tagged so the
auditor can treat it with appropriate skepticism — recognition errors on
code are common.

Legacy .ppt is not supported in v1; would require LibreOffice conversion.

Usage:
    python -m stale.tools.extract <file_or_directory>
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import fitz  # pymupdf
import pytesseract
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SUPPORTED_PDF = {".pdf"}
SUPPORTED_PPTX = {".pptx"}
LEGACY_PPT = {".ppt"}

OCR_MIN_CHARS = 8
OCR_PDF_DPI = 200


def extract_file(path: Path) -> str:
    """Extract text from a single file. Returns extracted text or a marker
    string for skipped/unsupported files. Catches malformed-file errors so
    one bad deck doesn't abort an entire course."""
    ext = path.suffix.lower()
    try:
        if ext in SUPPORTED_PDF:
            return _extract_pdf(path)
        if ext in SUPPORTED_PPTX:
            return _extract_pptx(path)
    except Exception as e:
        return f"[SKIPPED unreadable file {path.name}: {type(e).__name__}: {e}]"
    if ext in LEGACY_PPT:
        return f"[SKIPPED legacy .ppt format (not supported in v1): {path.name}]"
    return f"[SKIPPED unsupported extension {ext}: {path.name}]"


def extract_course(course_dir: Path) -> str:
    """Extract all decks in a course directory, concatenated with file markers.

    Output format:
        === FILE: <filename> ===
        <extracted text>
        === FILE: <next filename> ===
        ...

    Files are processed in sorted order so output is deterministic.
    Hidden files (.DS_Store, .git*) are skipped.
    """
    parts: list[str] = []
    for f in sorted(course_dir.iterdir()):
        if f.name.startswith("."):
            continue
        if not f.is_file():
            continue
        text = extract_file(f)
        parts.append(f"=== FILE: {f.name} ===\n\n{text}")
    return "\n\n".join(parts)


def extract_curriculum(curriculum_dir: Path) -> dict[str, str]:
    """Walk a curriculum directory (one subdirectory per course) and extract
    each course as a single concatenated text blob.

    Returns: {course_name: extracted_text}
    """
    out: dict[str, str] = {}
    for course in sorted(curriculum_dir.iterdir()):
        if course.name.startswith("."):
            continue
        if not course.is_dir():
            continue
        out[course.name] = extract_course(course)
    return out


def _ocr_image_bytes(blob: bytes) -> str:
    """OCR a single image. Returns stripped text, or empty string on failure
    or if recognition is too short to be meaningful."""
    try:
        img = Image.open(io.BytesIO(blob))
    except Exception:
        return ""
    try:
        text = pytesseract.image_to_string(img).strip()
    except Exception:
        return ""
    if len(text) < OCR_MIN_CHARS:
        return ""
    return text


def _extract_pdf(path: Path) -> str:
    """Extract text from a PDF page-by-page. For pages that embed images
    (likely code screenshots or diagrams) we additionally render the full
    page and OCR it, since fitz's get_text() can't see rasterized content.
    """
    doc = fitz.open(str(path))
    pages = []
    for i, page in enumerate(doc, 1):
        text = page.get_text().strip()
        chunks: list[str] = []
        if text:
            chunks.append(text)

        if page.get_images(full=False):
            try:
                pix = page.get_pixmap(dpi=OCR_PDF_DPI)
                ocr = _ocr_image_bytes(pix.tobytes("png"))
            except Exception:
                ocr = ""
            if ocr and ocr not in (text or ""):
                chunks.append(f"[OCR from page image]\n{ocr}")

        if chunks:
            pages.append(f"--- Page {i} ---\n" + "\n".join(chunks))
    doc.close()
    return "\n\n".join(pages)


def _extract_pptx(path: Path) -> str:
    prs = Presentation(str(path))
    slides_out = []
    for i, slide in enumerate(prs.slides, 1):
        chunks = [f"--- Slide {i} ---"]

        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    chunks.append(txt)
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    blob = shape.image.blob
                except Exception:
                    continue
                ocr = _ocr_image_bytes(blob)
                if ocr:
                    chunks.append(f"[OCR from picture]\n{ocr}")

        if slide.has_notes_slide:
            notes_tf = getattr(slide.notes_slide, "notes_text_frame", None)
            if notes_tf is not None:
                notes = notes_tf.text.strip()
                if notes:
                    chunks.append(f"[Speaker notes: {notes}]")

        if len(chunks) > 1:  # has content beyond the slide marker
            slides_out.append("\n".join(chunks))

    return "\n\n".join(slides_out)


def main() -> int:
    if len(sys.argv) != 2:
        print("Usage: python -m stale.tools.extract <file_or_directory>", file=sys.stderr)
        return 1

    target = Path(sys.argv[1]).expanduser().resolve()
    if not target.exists():
        print(f"Path does not exist: {target}", file=sys.stderr)
        return 1

    if target.is_file():
        print(extract_file(target))
        return 0

    # Directory — could be a single course or a curriculum (courses-of-courses)
    children = [p for p in target.iterdir() if not p.name.startswith(".")]
    if any(p.is_dir() for p in children):
        # treat as curriculum
        for course_name, text in extract_curriculum(target).items():
            print(f"\n\n========================================")
            print(f"COURSE: {course_name}")
            print(f"========================================\n")
            print(text)
    else:
        # treat as single course
        print(extract_course(target))
    return 0


if __name__ == "__main__":
    sys.exit(main())
